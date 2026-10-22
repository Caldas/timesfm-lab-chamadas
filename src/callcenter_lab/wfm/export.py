"""Export results to an Excel workbook (one sheet per result) with tidy formatting."""

from __future__ import annotations

from io import BytesIO

import pandas as pd

from .config import StaffingConfig
from .scheduling import Roster
from .shrinkage import daily_summary
from .tables import (
    breaks_table,
    config_table,
    coverage_table,
    plan_table,
    range_table,
    scenarios_result_table,
    sensitivity_result_table,
    shifts_table,
)

NOTICE = "Gerado pelo TimesFM Lab (laboratório de estudo). Os valores de exemplo não são de nenhuma operação real."


def _summary_table(config, plan, roster, scenarios) -> pd.DataFrame:
    rows: list[tuple[str, object]] = [("Aviso", NOTICE)]
    if plan is not None:
        calls = plan["calls"].sum()
        rows += [
            ("Primeiro intervalo", plan.index[0].strftime("%Y-%m-%d %H:%M")),
            ("Último intervalo", plan.index[-1].strftime("%Y-%m-%d %H:%M")),
            ("Chamadas no período", round(float(calls), 1)),
            ("Pico de agentes", int(plan["agents"].max())),
            ("Pico de pessoas escaladas", int(plan["scheduled"].max())),
            ("Nível de serviço esperado (%)",
             round(float((plan["service_level_pct"] * plan["calls"]).sum() / calls), 1) if calls else "n/d"),
        ]
        if config is not None:
            daily = daily_summary(plan, config.interval_minutes, config.shift_hours)
            rows += [("Horas de agente", round(float(daily["agent_hours"].sum()), 1)),
                     ("Horas escaladas", round(float(daily["scheduled_hours"].sum()), 1)),
                     ("FTE (soma dos dias)", round(float(daily["fte"].sum()), 1))]
        if "cost" in plan:
            rows.append(("Custo do dimensionamento", round(float(plan["cost"].sum()), 2)))
    if roster is not None:
        s = roster.summary()
        rows += [
            ("Escala: pessoas necessárias (headcount)", s["headcount"]),
            ("Escala: custo", round(s["cost"], 2)),
            ("Escala: intervalos com falta", s["shortfall_intervals"]),
            ("Escala: horas de falta", round(s["shortfall_person_hours"], 1)),
            ("Escala: horas de sobra", round(s["surplus_person_hours"], 1)),
            ("Escala: status do solver", roster.status),
        ]
    if scenarios is not None:
        rows.append(("Cenários comparados", len(scenarios)))
    return pd.DataFrame(rows, columns=["Item", "Valor"])


def build_workbook(
    *,
    config: StaffingConfig | None = None,
    plan: pd.DataFrame | None = None,
    plan_range: pd.DataFrame | None = None,
    roster: Roster | None = None,
    scenarios: pd.DataFrame | None = None,
    sensitivity: pd.DataFrame | None = None,
) -> bytes:
    """An .xlsx with a sheet for the summary, parameters, staffing plan, volume range, roster, scenarios and
    sensitivity: only the ones that were given. Returns the file content."""
    from openpyxl.styles import Alignment, Font, PatternFill

    sheets: list[tuple[str, pd.DataFrame, bool]] = [("Resumo", _summary_table(config, plan, roster, scenarios), False)]
    if config is not None:
        sheets.append(("Configuração", config_table(config), False))
    if plan is not None:
        sheets.append(("Dimensionamento", plan_table(plan).rename_axis("Intervalo"), True))
        if config is not None:
            daily = daily_summary(plan, config.interval_minutes, config.shift_hours).rename(columns={
                "agent_hours": "Horas de agente", "scheduled_hours": "Horas escaladas", "peak_agents": "Pico de agentes",
                "peak_scheduled": "Pico de pessoas escaladas", "fte": "FTE"}).round(1)
            daily.index = daily.index.strftime("%Y-%m-%d")
            sheets.append(("Por dia", daily.rename_axis("Dia"), True))
    if plan_range is not None:
        sheets.append(("Faixa de volume", range_table(plan_range).rename_axis("Intervalo"), True))
    if roster is not None:
        sheets += [("Escala - turnos", shifts_table(roster), False), ("Escala - pausas", breaks_table(roster), False),
                   ("Escala - cobertura", coverage_table(roster).rename_axis("Intervalo"), True)]
    if scenarios is not None:
        sheets.append(("Cenários", scenarios_result_table(scenarios), False))
    if sensitivity is not None:
        sheets.append(("Sensibilidade", sensitivity_result_table(sensitivity), False))

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm") as writer:
        for name, frame, with_index in sheets:
            frame.to_excel(writer, sheet_name=name[:31], index=with_index)
            sheet = writer.sheets[name[:31]]
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E79")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            sheet.freeze_panes = "A2"
            for column in sheet.columns:
                longest = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
                sheet.column_dimensions[column[0].column_letter].width = min(max(longest + 2, 10), 60)
    return buffer.getvalue()
